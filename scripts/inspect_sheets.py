#!/usr/bin/env python3
"""
Inspect sheets in itinerary.xlsx:
Outputs sheet names, header columns, and row counts.
"""
import sys
from pathlib import Path
import openpyxl

def inspect_workbook(file_path: str = "itinerary.xlsx"):
    path = Path(file_path)
    if not path.exists():
        print(f"Error: File {file_path} not found.", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(path, data_only=True)
    print("=" * 60)
    print(f"Workbook: {path.name} (Sheets: {len(wb.sheetnames)})")
    print("=" * 60)

    for idx, sheet_name in enumerate(wb.sheetnames, 1):
        sheet = wb[sheet_name]
        total_rows = sheet.max_row or 0
        total_cols = sheet.max_column or 0
        
        # Find first non-empty row as header
        headers = []
        header_row_idx = None
        data_row_count = 0

        for r_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
            non_empty_cells = [c for c in row if c is not None and str(c).strip() != ""]
            if non_empty_cells:
                if header_row_idx is None:
                    header_row_idx = r_idx
                    headers = [str(c).strip() if c is not None else f"Col_{i+1}" for i, c in enumerate(row)]
                    # Trim trailing empty header columns
                    while headers and headers[-1].startswith("Col_"):
                        headers.pop()
                else:
                    data_row_count += 1

        print(f"\n[{idx}] Sheet: {sheet_name}")
        print(f"    - Dimensions: {total_rows} rows x {total_cols} cols")
        print(f"    - Header Row: Line {header_row_idx if header_row_idx else 'None'}")
        print(f"    - Headers: {headers}")
        print(f"    - Valid Data Rows: {data_row_count}")

    print("\n" + "=" * 60)
    print("Inspection completed.")
    print("=" * 60)

if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "itinerary.xlsx"
    inspect_workbook(target)
